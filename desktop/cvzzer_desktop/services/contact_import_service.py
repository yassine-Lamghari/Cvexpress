from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from ..models import Contact


EMAIL_PATTERN = re.compile(r'^[^\s@]+@[^\s@]+\.[^\s@]+$')
HEADER_ALIASES = {
    'email': {'email', 'e-mail', 'mail', 'adresse email', 'adresse e-mail'},
    'company': {'company', 'entreprise', 'societe', 'société'},
    'contact_name': {'contact_name', 'contact', 'nom', 'nom contact', 'recruteur'},
    'position': {'position', 'poste', 'role', 'rôle', 'fonction'},
}


@dataclass
class ContactImportResult:
    contacts: list[Contact] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    duplicates: int = 0


def _normalise(value: object) -> str:
    return str(value or '').strip()


def _field_map(headers: list[str]) -> dict[str, int]:
    normalised = [header.lower().strip() for header in headers]
    mapping: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for index, header in enumerate(normalised):
            if header in aliases:
                mapping[field] = index
                break

    if 'email' not in mapping:
        for index, header in enumerate(normalised):
            if 'mail' in header:
                mapping['email'] = index
                break
    return mapping


def _parse_rows(rows: list[list[object]]) -> ContactImportResult:
    result = ContactImportResult()
    if not rows:
        result.errors.append('Le fichier ne contient aucune ligne.')
        return result

    headers = [_normalise(value) for value in rows[0]]
    mapping = _field_map(headers)
    if 'email' not in mapping:
        result.errors.append('Aucune colonne email détectée. Utilisez par exemple « email » ou « e-mail ».')
        return result

    seen: set[str] = set()
    for row_number, row in enumerate(rows[1:], start=2):
        def read(field: str) -> str:
            index = mapping.get(field)
            return _normalise(row[index]) if index is not None and index < len(row) else ''

        email = read('email').lower()
        if not email:
            continue
        if not EMAIL_PATTERN.fullmatch(email):
            result.errors.append(f'Ligne {row_number} : adresse email invalide ({email}).')
            continue
        if email in seen:
            result.duplicates += 1
            continue
        seen.add(email)
        result.contacts.append(Contact(email, read('company'), read('contact_name'), read('position')))
    return result


def import_contacts(path: str | Path) -> ContactImportResult:
    source = Path(path)
    if source.suffix.lower() == '.csv':
        with source.open('r', encoding='utf-8-sig', newline='') as stream:
            return _parse_rows([list(row) for row in csv.reader(stream)])
    if source.suffix.lower() not in {'.xlsx', '.xlsm'}:
        return ContactImportResult(errors=['Format non pris en charge. Utilisez un fichier .xlsx, .xlsm ou .csv.'])

    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook.active
    return _parse_rows([list(row) for row in worksheet.iter_rows(values_only=True)])
